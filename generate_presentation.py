import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    # 16:9 Widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6] # completely blank layout

    # Color Palette - Ninja Dojo / Wood & Gold Theme
    BG_COLOR = RGBColor(24, 16, 11)          # Deep Mahogany Dojo Wood
    CARD_BG = RGBColor(40, 27, 18)           # Rich Dark Wood Plank Card
    CARD_BORDER = RGBColor(160, 115, 55)     # Warm Golden Ochre
    CARD_HEADER_BG = RGBColor(56, 38, 25)    # Slightly lighter wood for headers
    
    TEXT_GOLD = RGBColor(250, 195, 75)       # Bright Ninja Gold
    TEXT_CREAM = RGBColor(245, 240, 230)     # Clean Cream White
    TEXT_MUTED = RGBColor(185, 170, 150)     # Tan Khaki Muted
    ACCENT_RED = RGBColor(220, 50, 35)       # Martial Crimson
    ACCENT_GREEN = RGBColor(76, 175, 80)     # Ninja Jade / Kiwi Green
    ACCENT_ORANGE = RGBColor(235, 125, 25)   # Amber Orange

    SPRITE_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Assets", "Sprites")
    
    def get_sprite(filename):
        path = os.path.join(SPRITE_DIR, filename)
        return path if os.path.exists(path) else None

    watermelon_img = get_sprite("Dưa hấu.png")
    apple_img = get_sprite("táo.png")
    orange_img = get_sprite("Cam.png")
    banana_img = get_sprite("Chuoi.png")
    strawberry_img = get_sprite("DauTay.png")
    bomb_img = get_sprite("bomb.png")

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5)
        )
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_COLOR
        bg.line.fill.background() # no line

        # Subtle decorative top accent bar
        bar = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.08)
        )
        bar.fill.solid()
        bar.fill.fore_color.rgb = TEXT_GOLD
        bar.line.fill.background()

    def add_header(slide, badge_text, title_text, subtitle_text=None):
        # Category Badge
        badge_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(8), Inches(0.35))
        tf_badge = badge_box.text_frame
        tf_badge.word_wrap = True
        tf_badge.margin_left = tf_badge.margin_top = tf_badge.margin_right = tf_badge.margin_bottom = 0
        p_badge = tf_badge.paragraphs[0]
        p_badge.text = badge_text.upper()
        p_badge.font.size = Pt(11)
        p_badge.font.bold = True
        p_badge.font.color.rgb = TEXT_GOLD
        p_badge.font.name = "Segoe UI"

        # Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.72), Inches(10.5), Inches(0.7))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        tf_title.margin_left = tf_title.margin_top = tf_title.margin_right = tf_title.margin_bottom = 0
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(26)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_CREAM
        p_title.font.name = "Segoe UI"

        if subtitle_text:
            p_title.font.size = Pt(24)
            sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.35), Inches(10.5), Inches(0.35))
            tf_sub = sub_box.text_frame
            tf_sub.word_wrap = True
            tf_sub.margin_left = tf_sub.margin_top = tf_sub.margin_right = tf_sub.margin_bottom = 0
            p_sub = tf_sub.paragraphs[0]
            p_sub.text = subtitle_text
            p_sub.font.size = Pt(13)
            p_sub.font.color.rgb = TEXT_MUTED
            p_sub.font.name = "Segoe UI"

    def create_card(slide, left, top, width, height, border_color=CARD_BORDER, bg_color=CARD_BG):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1.2)
        return card

    # ==========================================
    # SLIDE 1: TITLE SLIDE
    # ==========================================
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1)

    # Decorative frame in center
    frame = create_card(slide1, Inches(0.8), Inches(0.8), Inches(11.733), Inches(5.9), CARD_BORDER, CARD_BG)

    # Top Tag
    tag_box = slide1.shapes.add_textbox(Inches(1.2), Inches(1.2), Inches(9.0), Inches(0.4))
    p_tag = tag_box.text_frame.paragraphs[0]
    p_tag.text = "★  ĐỒ ÁN LẬP TRÌNH GAME 2D  •  CÔNG NGHỆ THÔNG TIN  ★"
    p_tag.font.size = Pt(12)
    p_tag.font.bold = True
    p_tag.font.color.rgb = TEXT_GOLD
    p_tag.font.name = "Segoe UI"

    # Main Project Title
    t_box = slide1.shapes.add_textbox(Inches(1.2), Inches(1.65), Inches(9.5), Inches(1.2))
    p_t = t_box.text_frame.paragraphs[0]
    p_t.text = "FRUIT NINJA 2D"
    p_t.font.size = Pt(46)
    p_t.font.bold = True
    p_t.font.color.rgb = TEXT_GOLD
    p_t.font.name = "Segoe UI"

    # Subtitle
    sub_box = slide1.shapes.add_textbox(Inches(1.2), Inches(2.8), Inches(9.5), Inches(0.6))
    p_sub = sub_box.text_frame.paragraphs[0]
    p_sub.text = "Game Chém Hoa Quả Phong Cách Võ Quán Cổ Điển"
    p_sub.font.size = Pt(20)
    p_sub.font.bold = True
    p_sub.font.color.rgb = TEXT_CREAM
    p_sub.font.name = "Segoe UI"

    # Divider line
    div = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.2), Inches(3.5), Inches(7.5), Inches(0.04))
    div.fill.solid()
    div.fill.fore_color.rgb = ACCENT_ORANGE
    div.line.fill.background()

    # Meta Info Cards
    info_card = create_card(slide1, Inches(1.2), Inches(3.8), Inches(6.8), Inches(2.2), RGBColor(100, 75, 45), RGBColor(32, 22, 15))
    tb_info = slide1.shapes.add_textbox(Inches(1.4), Inches(3.95), Inches(6.4), Inches(1.9))
    tf_info = tb_info.text_frame
    tf_info.word_wrap = True

    infos = [
        ("Giảng viên hướng dẫn:", "ThS. [Họ và Tên Giảng Viên]"),
        ("Sinh viên thực hiện:", "[Họ và Tên Sinh Viên]"),
        ("Mã số sinh viên (MSSV):", "[2xxxxxxx]"),
        ("Nền tảng & Công nghệ:", "Unity Engine (URP 2D)  •  C# Scripting")
    ]
    for i, (label, val) in enumerate(infos):
        p = tf_info.paragraphs[0] if i == 0 else tf_info.add_paragraph()
        r1 = p.add_run()
        r1.text = f"• {label} "
        r1.font.bold = True
        r1.font.size = Pt(13)
        r1.font.color.rgb = TEXT_GOLD
        r1.font.name = "Segoe UI"

        r2 = p.add_run()
        r2.text = val
        r2.font.bold = False
        r2.font.size = Pt(13)
        r2.font.color.rgb = TEXT_CREAM
        r2.font.name = "Segoe UI"
        p.space_after = Pt(6)

    # Fruit decorative illustration on right
    if watermelon_img:
        slide1.shapes.add_picture(watermelon_img, Inches(8.5), Inches(1.8), width=Inches(2.5))
    if bomb_img:
        slide1.shapes.add_picture(bomb_img, Inches(10.2), Inches(3.6), width=Inches(1.8))
    if orange_img:
        slide1.shapes.add_picture(orange_img, Inches(8.3), Inches(4.3), width=Inches(1.6))

    # ==========================================
    # SLIDE 2: LÝ DO CHỌN ĐỀ TÀI & MỤC TIÊU
    # ==========================================
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide2)
    add_header(slide2, "Tổng quan dự án", "LÝ DO CHỌN ĐỀ TÀI & MỤC TIÊU PHÁT TRIỂN", "Định hướng nghiên cứu và hiện thực hóa sản phẩm game 2D tương tác trực quan")

    cards_data_s2 = [
        ("01", "TÍNH CẤP THIẾT & THỰC TIỄN", TEXT_GOLD, [
            ("Lối chơi kinh điển:", "Fruit Ninja là tựa game casual có sức hút toàn cầu, tính giải trí cao, dễ tiếp cận với mọi đối tượng."),
            ("Tương tác tức thì:", "Đòi hỏi phản xạ nhanh, trải nghiệm xúc giác sảng khoái với vệt chém kiếm (Blade Slash)."),
            ("Giá trị đồ án:", "Phù hợp hoàn hảo với khối lượng kiến thức đồ án chuyên ngành CNTT / Game Development.")
        ]),
        ("02", "MỤC TIÊU KỸ THUẬT", ACCENT_ORANGE, [
            ("Làm chủ Unity 2D:", "Vận dụng thành thạo Rigidbody2D, CircleCollider2D và hệ thống Raycasting."),
            ("Tối ưu hóa bộ nhớ:", "Hiện thực giải pháp Object Pooling để triệt tiêu hiện tượng sụt FPS do Garbage Collector."),
            ("Mẫu thiết kế chuẩn:", "Ứng dụng mô hình Singleton, Event-driven và Component-based trong mã nguồn C#.")
        ]),
        ("03", "MỤC TIÊU SẢN PHẨM", ACCENT_GREEN, [
            ("Mượt mà 60 FPS:", "Game vận hành ổn định trên cả máy tính cấu hình phổ thông lẫn thiết bị cảm ứng."),
            ("Hiệu ứng chân thực:", "Âm thanh chém, tiếng nổ bom, hiệu ứng té nước ép và rung màn hình (Camera Shake)."),
            ("Khả năng mở rộng:", "Dễ dàng bổ sung thêm trái cây đặc biệt, hiệu ứng kiếm và chế độ chơi mới.")
        ]),
    ]

    card_w = Inches(3.64)
    card_h = Inches(5.1)
    for idx, (num, title, accent_c, points) in enumerate(cards_data_s2):
        left_pos = Inches(0.8 + idx * 4.04)
        create_card(slide2, left_pos, Inches(1.8), card_w, card_h)

        # Header bar inside card
        header_bar = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_pos + Inches(0.15), Inches(1.95), card_w - Inches(0.3), Inches(0.8))
        header_bar.fill.solid()
        header_bar.fill.fore_color.rgb = CARD_HEADER_BG
        header_bar.line.color.rgb = accent_c
        header_bar.line.width = Pt(1)

        # Number badge & title
        tb = slide2.shapes.add_textbox(left_pos + Inches(0.2), Inches(2.0), card_w - Inches(0.4), Inches(0.7))
        tf = tb.text_frame
        tf.word_wrap = True
        p_num = tf.paragraphs[0]
        p_num.text = f"MỤC TIÊU {num}  •  {title}"
        p_num.font.bold = True
        p_num.font.size = Pt(12)
        p_num.font.color.rgb = accent_c
        p_num.font.name = "Segoe UI"

        # Content bullets
        tb_body = slide2.shapes.add_textbox(left_pos + Inches(0.2), Inches(2.9), card_w - Inches(0.4), Inches(3.8))
        tf_body = tb_body.text_frame
        tf_body.word_wrap = True
        for p_idx, (bold_txt, norm_txt) in enumerate(points):
            p = tf_body.paragraphs[0] if p_idx == 0 else tf_body.add_paragraph()
            r_b = p.add_run()
            r_b.text = f"✔ {bold_txt} "
            r_b.font.bold = True
            r_b.font.size = Pt(12)
            r_b.font.color.rgb = TEXT_CREAM
            r_b.font.name = "Segoe UI"

            r_n = p.add_run()
            r_n.text = norm_txt
            r_n.font.bold = False
            r_n.font.size = Pt(11.5)
            r_n.font.color.rgb = TEXT_MUTED
            r_n.font.name = "Segoe UI"
            p.space_after = Pt(12)

    # ==========================================
    # SLIDE 3: TỔNG QUAN GAMEPLAY & LUẬT CHƠI
    # ==========================================
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide3)
    add_header(slide3, "Cơ chế trò chơi", "TỔNG QUAN GAMEPLAY & LUẬT CHƠI", "Lối chơi trực quan, tiết tấu nhanh, đề cao sự tập trung và phản xạ chính xác")

    gameplay_cards = [
        ("⚔️  VỆT KIẾM CHÉM (BLADE SLICE)", TEXT_GOLD, [
            ("Thao tác điều khiển:", "Người chơi nhấn giữ chuột hoặc chạm ngón tay vuốt ngang màn hình."),
            ("Điều kiện chém hợp lệ:", "Vận tốc di chuyển phải vượt ngưỡng Min Velocity để kích hoạt đường kiếm sắc bén."),
            ("Độ chính xác cao:", "Vệt sáng TrailRenderer bám sát thao tác, tạo cảm giác nhập vai kiếm khách chân thực.")
        ]),
        ("🍉  QUẢ NẢY VẬT LÝ TỰ NHIÊN", ACCENT_GREEN, [
            ("Quỹ đạo Parabol:", "Trái cây bắn lên ngẫu nhiên từ đáy màn hình với góc nảy và lực văng ngẫu nhiên."),
            ("Đa dạng chủng loại:", "Táo, Chuối, Cam, Dâu tây, Dưa hấu - mỗi loại mang màu sắc và kích thước riêng biệt."),
            ("Tách đôi đẹp mắt:", "Khi bị chém trúng, quả tách đôi thành 2 nửa FruitHalf văng xoay sang hai bên.")
        ]),
        ("💣  CHƯỚNG NGẠI VẬT & BẪY BOM", ACCENT_RED, [
            ("Yếu tố bất ngờ:", "Bom đen lẫn vào giữa các đợt bắn trái cây để thử thách độ tinh mắt của người chơi."),
            ("Hậu quả nghiêm trọng:", "Chém trúng bom lập tức kích hoạt nổ, rung màn hình và bị trừ 1 mạng sống."),
            ("Chiến thuật né tránh:", "Đòi hỏi người chơi kiểm soát đường kiếm dừng đúng lúc để không chạm trúng bom.")
        ]),
        ("❤️  MẠNG SỐNG & ĐỘ KHÓ LŨY TIẾN", ACCENT_ORANGE, [
            ("Khởi đầu với 3 Tim:", "Người chơi sở hữu 3 mạng sống; mất hết 3 mạng trò chơi sẽ kết thúc (Game Over)."),
            ("Lối chơi xả stress:", "Quả rơi xuống màn hình không bị trừ mạng, giúp game thủ tập trung vào việc tạo combo."),
            ("Độ khó tăng liên tục:", "Mỗi 10 giây, tần suất bắn và số lượng quả tăng lên 1.15 lần, đẩy nhịp độ lên cao trào.")
        ])
    ]

    for idx, (title, accent_c, pts) in enumerate(gameplay_cards):
        col = idx % 2
        row = idx // 2
        l = Inches(0.8 + col * 5.95)
        t = Inches(1.8 + row * 2.65)
        w = Inches(5.75)
        h = Inches(2.45)
        create_card(slide3, l, t, w, h)

        tb_head = slide3.shapes.add_textbox(l + Inches(0.2), t + Inches(0.12), w - Inches(0.4), Inches(0.45))
        p_h = tb_head.text_frame.paragraphs[0]
        p_h.text = title
        p_h.font.bold = True
        p_h.font.size = Pt(14)
        p_h.font.color.rgb = accent_c
        p_h.font.name = "Segoe UI"

        tb_c = slide3.shapes.add_textbox(l + Inches(0.2), t + Inches(0.55), w - Inches(0.4), Inches(1.8))
        tf_c = tb_c.text_frame
        tf_c.word_wrap = True
        for p_i, (bld, nrm) in enumerate(pts):
            p = tf_c.paragraphs[0] if p_i == 0 else tf_c.add_paragraph()
            rb = p.add_run()
            rb.text = f"• {bld} "
            rb.font.bold = True
            rb.font.size = Pt(11.5)
            rb.font.color.rgb = TEXT_CREAM
            rb.font.name = "Segoe UI"

            rn = p.add_run()
            rn.text = nrm
            rn.font.bold = False
            rn.font.size = Pt(11)
            rn.font.color.rgb = TEXT_MUTED
            rn.font.name = "Segoe UI"
            p.space_after = Pt(4)

    # ==========================================
    # SLIDE 4: CÔNG NGHỆ & CÔNG CỤ PHÁT TRIỂN
    # ==========================================
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide4)
    add_header(slide4, "Công nghệ ứng dụng", "NGĂN XẾP CÔNG NGHỆ & CÔNG CỤ PHÁT TRIỂN", "Lựa chọn công nghệ chuẩn ngành game, đảm bảo tính ổn định và khả năng tối ưu cao")

    tech_data = [
        ("UNITY ENGINE", "2022 / 6 LTS (URP 2D)", TEXT_GOLD, [
            "Công cụ làm game hàng đầu thế giới cho thể loại 2D/3D.",
            "Universal Render Pipeline (URP 2D) tối ưu ánh sáng & đồ họa.",
            "Tương thích đa nền tảng: PC Windows, WebGL, Android, iOS."
        ]),
        ("NGÔN NGỮ C#", "Object-Oriented Programming", ACCENT_ORANGE, [
            "Lập trình toàn bộ Gameplay logic, Physics & Event handling.",
            "Tận dụng Generic Collections (Queue, List, HashSet) tối ưu tốc độ.",
            "Tách biệt rõ ràng các phân hệ xử lý bằng Design Patterns."
        ]),
        ("VẬT LÝ UNITY 2D", "Physics2D & Collision Matrix", ACCENT_GREEN, [
            "Rigidbody2D: Mô phỏng trọng lực g=-9.8m/s² và quỹ đạo Parabol chân thực.",
            "Raycast2D: Tính toán đường cắt kiếm siêu tốc không độ trễ.",
            "CircleCollider2D & LayerMask: Phân tầng Fruit / Bomb độc lập."
        ]),
        ("GRAPHICS & AUDIO", "Assets & Post-Processing", ACCENT_RED, [
            "Sprite 2D & TextMeshPro: Hiển thị font chữ và UI sắc nét chuẩn HD.",
            "TrailRenderer: Vẽ vệt kiếm phát sáng chuyển động mượt mà.",
            "Unity Audio Source: Hiệu ứng âm thanh chém kiếm & bom nổ sống động."
        ])
    ]

    card_w4 = Inches(2.76)
    card_h4 = Inches(5.1)
    for idx, (tech_title, tech_sub, accent_c, bullet_list) in enumerate(tech_data):
        l = Inches(0.8 + idx * 2.99)
        create_card(slide4, l, Inches(1.8), card_w4, card_h4)

        # Header Box
        hb = slide4.shapes.add_shape(MSO_SHAPE.RECTANGLE, l + Inches(0.12), Inches(1.95), card_w4 - Inches(0.24), Inches(0.95))
        hb.fill.solid()
        hb.fill.fore_color.rgb = CARD_HEADER_BG
        hb.line.color.rgb = accent_c
        hb.line.width = Pt(1)

        tb_th = slide4.shapes.add_textbox(l + Inches(0.15), Inches(2.0), card_w4 - Inches(0.3), Inches(0.9))
        p_th = tb_th.text_frame.paragraphs[0]
        p_th.text = tech_title
        p_th.font.bold = True
        p_th.font.size = Pt(13)
        p_th.font.color.rgb = accent_c
        p_th.font.name = "Segoe UI"

        p_ts = tb_th.text_frame.add_paragraph()
        p_ts.text = tech_sub
        p_ts.font.size = Pt(10)
        p_ts.font.color.rgb = TEXT_MUTED
        p_ts.font.name = "Segoe UI"

        # Content
        tb_tc = slide4.shapes.add_textbox(l + Inches(0.15), Inches(3.05), card_w4 - Inches(0.3), Inches(3.7))
        tf_tc = tb_tc.text_frame
        tf_tc.word_wrap = True
        for b_i, b_item in enumerate(bullet_list):
            p = tf_tc.paragraphs[0] if b_i == 0 else tf_tc.add_paragraph()
            p.text = f"❖ {b_item}"
            p.font.size = Pt(11)
            p.font.color.rgb = TEXT_CREAM
            p.font.name = "Segoe UI"
            p.space_after = Pt(12)

    # ==========================================
    # SLIDE 5: KIẾN TRÚC HỆ THỐNG
    # ==========================================
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide5)
    add_header(slide5, "Thiết kế phần mềm", "KIẾN TRÚC HỆ THỐNG & CÁC MODULE CHÍNH", "Tổ chức mã nguồn theo mô hình Component-Based kết hợp Singleton Pattern vững chắc")

    # Left box: Architecture Diagram Concept
    create_card(slide5, Inches(0.8), Inches(1.8), Inches(4.5), Inches(5.1), TEXT_GOLD)
    tb_diag_title = slide5.shapes.add_textbox(Inches(1.0), Inches(1.95), Inches(4.1), Inches(0.4))
    p_dt = tb_diag_title.text_frame.paragraphs[0]
    p_dt.text = "MÔ HÌNH QUẢN LÝ TẬP TRUNG"
    p_dt.font.bold = True
    p_dt.font.size = Pt(14)
    p_dt.font.color.rgb = TEXT_GOLD

    modules_arch = [
        ("GameManager", "Quản lý Game State, Mạng sống, Bộ đếm độ khó", TEXT_GOLD),
        ("ScoreManager", "Lưu trữ Điểm số, Combo Streak, High Score", ACCENT_GREEN),
        ("UIManager", "Cập nhật HUD, Màn hình Menu, Paused, Game Over", ACCENT_ORANGE),
        ("FruitSpawner", "Tính toán nhịp độ bắn, tạo góc nảy & lực đẩy", TEXT_CREAM),
        ("ObjectPool", "Kho lưu trữ tái sử dụng Fruit, HalfFruit & Bomb", ACCENT_RED),
        ("BladeController", "Bắt sự kiện vuốt chuột, tính vận tốc & Raycast", TEXT_GOLD)
    ]
    for m_idx, (m_name, m_desc, m_col) in enumerate(modules_arch):
        m_top = Inches(2.45 + m_idx * 0.72)
        mb = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), m_top, Inches(4.1), Inches(0.62))
        mb.fill.solid()
        mb.fill.fore_color.rgb = CARD_HEADER_BG
        mb.line.color.rgb = m_col
        mb.line.width = Pt(1)

        tb_m = slide5.shapes.add_textbox(Inches(1.1), m_top + Inches(0.04), Inches(3.9), Inches(0.55))
        tf_m = tb_m.text_frame
        p1 = tf_m.paragraphs[0]
        p1.text = f"★ {m_name}"
        p1.font.bold = True
        p1.font.size = Pt(12)
        p1.font.color.rgb = m_col

        p2 = tf_m.add_paragraph()
        p2.text = m_desc
        p2.font.size = Pt(9.5)
        p2.font.color.rgb = TEXT_MUTED

    # Right side: 3 Deep Dive Cards
    right_cards = [
        ("MẪU THIẾT KẾ SINGLETON (SINGLETON PATTERN)", TEXT_GOLD, [
            ("Cơ chế hoạt động:", "Đảm bảo mỗi Manager (GameManager, ScoreManager, UIManager) chỉ có 1 thực thể duy nhất tồn tại."),
            ("Ưu điểm vượt trội:", "Các script khác dễ dàng truy cập thông qua `GameManager.Instance` mà không cần tham chiếu kéo thả thủ công trên Inspector."),
            ("Bảo toàn dữ liệu:", "Duy trì trạng thái xuyên suốt quá trình chơi mà không bị mất dữ liệu giữa các lần Reset.")
        ]),
        ("GIAO TIẾP EVENT-DRIVEN & LOOSE COUPLING", ACCENT_ORANGE, [
            ("Tách biệt trách nhiệm (Separation of Concerns):", "Logic tính điểm hoàn toàn tách rời logic hiển thị giao diện UI."),
            ("Phản xạ thời gian thực:", "Khi quả bị chém, Fruit gọi `ScoreManager.AddScore()`, sau đó ScoreManager tự động báo cho `UIManager` cập nhật TextMeshPro.")
        ]),
        ("TỐI ƯU HÓA LUỒNG XỬ LÝ (GAME LOOP)", ACCENT_GREEN, [
            ("Vòng lặp Update tinh gọn:", "Chỉ xử lý logic chém và spawn khi State đang ở `GameState.Playing`."),
            ("Kiểm soát thời gian (TimeScale):", "Dừng game tức thì (`Time.timeScale = 0`) khi Pause hoặc Game Over mà không gây lỗi luồng vật lý.")
        ])
    ]

    for idx, (rt_title, rt_col, rt_pts) in enumerate(right_cards):
        top_pos = Inches(1.8 + idx * 1.73)
        create_card(slide5, Inches(5.55), top_pos, Inches(6.98), Inches(1.6))

        tb_rt = slide5.shapes.add_textbox(Inches(5.75), top_pos + Inches(0.08), Inches(6.6), Inches(0.35))
        p_rh = tb_rt.text_frame.paragraphs[0]
        p_rh.text = rt_title
        p_rh.font.bold = True
        p_rh.font.size = Pt(12.5)
        p_rh.font.color.rgb = rt_col

        tb_rb = slide5.shapes.add_textbox(Inches(5.75), top_pos + Inches(0.38), Inches(6.6), Inches(1.15))
        tf_rb = tb_rb.text_frame
        tf_rb.word_wrap = True
        for pt_i, (bld, nrm) in enumerate(rt_pts):
            p = tf_rb.paragraphs[0] if pt_i == 0 else tf_rb.add_paragraph()
            rb = p.add_run()
            rb.text = f"✔ {bld} "
            rb.font.bold = True
            rb.font.size = Pt(10.5)
            rb.font.color.rgb = TEXT_CREAM

            rn = p.add_run()
            rn.text = nrm
            rn.font.bold = False
            rn.font.size = Pt(10.5)
            rn.font.color.rgb = TEXT_MUTED
            p.space_after = Pt(2)

    # ==========================================
    # SLIDE 6: KỸ THUẬT CỐT LÕI: VẾT CHÉM & CẮT QUẢ
    # ==========================================
    slide6 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide6)
    add_header(slide6, "Thuật toán xử lý", "KỸ THUẬT CỐT LÕI: VẾT CHÉM & PHÂN TÁCH TRÁI CÂY", "Kết hợp mượt mà giữa tính toán hình học Vector2, Raycast2D và hiệu ứng vật lý Rigidbody")

    slicing_cards = [
        ("01. NHẬN DIỆN VẾT CHÉM (RAYCAST 2D)", TEXT_GOLD, [
            ("Thuật toán bám đuổi tọa độ:", "Liên tục lưu lại vị trí chuột ở frame trước (`previousWorldPos`) và frame hiện tại (`currentWorldPos`)."),
            ("Ngưỡng vận tốc (Min Velocity):", "Tính `velocity = distance / deltaTime`. Nếu vận tốc nhỏ hơn ngưỡng 0.5m/s thì không tính là nhát chém hợp lệ."),
            ("Raycasting chính xác 100%:", "Bắn tia `Physics2D.Linecast()` giữa 2 tọa độ. Đảm bảo dù chuột quẹt nhanh đến đâu cũng không bị lọt qua khe va chạm."),
            ("Bộ lọc HashSet (Anti-duplicate):", "Sử dụng `HashSet<EntityId>` để đảm bảo mỗi quả chỉ bị chém trúng 1 lần trong 1 nhát quẹt.")
        ]),
        ("02. CƠ CHẾ CẮT ĐÔI QUẢ (FRUIT SPLITTING)", ACCENT_ORANGE, [
            ("Thay thế tức thời:", "Ngay khi va chạm xảy ra, quả nguyên lập tức ẩn đi (`SetActive: false`) và trả về Pool."),
            ("Kích hoạt 2 nửa FruitHalf:", "Lấy 2 nửa quả đã được gán sẵn Sprite bán nguyệt tương ứng từ Object Pool."),
            ("Gia tốc vật lý trái chiều:", "Gán vận tốc tức thời: Nửa trái văng sang góc trái (-X), nửa phải văng sang góc phải (+X)."),
            ("Hiệu ứng xoay ngẫu nhiên:", "Áp dụng `AddTorque(randomSpin)` khiến 2 nửa quả xoay tròn khi rơi tự do xuống dưới.")
        ]),
        ("03. HIỆU ỨNG THỊ GIÁC & PHẢN HỒI XÚC GIÁC", ACCENT_GREEN, [
            ("TrailRenderer động:", "Vệt kiếm ninja phát sáng, tự động co nhỏ đuôi theo thời gian (`time = 0.15s`), tạo cảm giác bén ngọt."),
            ("Camera Shake (Rung chấn):", "Khi chém trúng bom hoặc combo lớn, camera rung giật trong 0.2s để tạo ấn tượng va đập cực mạnh."),
            ("Splash Particles:", "Hạt bụi nước ép trái cây tung tóe tại điểm chém, gia tăng sự phấn khích cho người chơi.")
        ])
    ]

    for idx, (title, col, pts) in enumerate(slicing_cards):
        l = Inches(0.8 + idx * 4.04)
        create_card(slide6, l, Inches(1.8), Inches(3.64), Inches(5.1))

        hb = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l + Inches(0.15), Inches(1.95), Inches(3.34), Inches(0.8))
        hb.fill.solid()
        hb.fill.fore_color.rgb = CARD_HEADER_BG
        hb.line.color.rgb = col
        hb.line.width = Pt(1)

        tb_h = slide6.shapes.add_textbox(l + Inches(0.2), Inches(2.0), Inches(3.24), Inches(0.7))
        p_h = tb_h.text_frame.paragraphs[0]
        p_h.text = title
        p_h.font.bold = True
        p_h.font.size = Pt(11.5)
        p_h.font.color.rgb = col

        tb_b = slide6.shapes.add_textbox(l + Inches(0.18), Inches(2.9), Inches(3.28), Inches(3.8))
        tf_b = tb_b.text_frame
        tf_b.word_wrap = True
        for p_i, (bld, nrm) in enumerate(pts):
            p = tf_b.paragraphs[0] if p_i == 0 else tf_b.add_paragraph()
            rb = p.add_run()
            rb.text = f"❖ {bld} "
            rb.font.bold = True
            rb.font.size = Pt(11)
            rb.font.color.rgb = TEXT_CREAM

            rn = p.add_run()
            rn.text = nrm
            rn.font.bold = False
            rn.font.size = Pt(10.5)
            rn.font.color.rgb = TEXT_MUTED
            p.space_after = Pt(8)

    # ==========================================
    # SLIDE 7: HỆ THỐNG SPAWNER & TỐI ƯU OBJECT POOLING
    # ==========================================
    slide7 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide7)
    add_header(slide7, "Tối ưu hóa hiệu năng", "HỆ THỐNG SPAWNER & GIẢI PHÁP OBJECT POOLING", "Đảm bảo game luôn duy trì 60+ FPS ổn định, triệt tiêu tình trạng nghẽn bộ nhớ")

    # Left Card: Fruit Spawner
    create_card(slide7, Inches(0.8), Inches(1.8), Inches(5.75), Inches(5.1), TEXT_GOLD)
    tb_sp_title = slide7.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(5.15), Inches(0.5))
    p_st = tb_sp_title.text_frame.paragraphs[0]
    p_st.text = "🎯 THUẬT TOÁN BẮN QUẢ (FRUIT SPAWNER)"
    p_st.font.bold = True
    p_st.font.size = Pt(14)
    p_st.font.color.rgb = TEXT_GOLD

    spawner_points = [
        ("Vị trí bắn ngẫu nhiên:", "Trái cây xuất hiện dọc mép dưới màn hình (X từ -ScreenBound đến +ScreenBound)."),
        ("Góc bắn thông minh:", "Góc nảy luôn hướng về tâm màn hình (Center Screen), tránh việc quả bị văng ra khỏi tầm nhìn của người chơi."),
        ("Lực nảy Parabol ngẫu nhiên:", "Áp dụng lực `AddForce(Vector2 * randomForce)` để chiều cao đỉnh rơi luôn đa dạng, tạo thử thách phong phú."),
        ("Bộ điều phối đợt bắn (Waves):", "Xen kẽ giữa bắn quả đơn lẻ, bắn theo chùm (Multi-fruit) và bất ngờ lồng ghép bom nguy hiểm."),
        ("Hệ số khó lũy tiến (Difficulty Factor):", "Cứ mỗi 10 giây, chu kỳ nghỉ giữa các đợt bắn giảm dần, số lượng quả tăng lên.")
    ]
    tb_sp_body = slide7.shapes.add_textbox(Inches(1.1), Inches(2.6), Inches(5.15), Inches(4.1))
    tf_sp = tb_sp_body.text_frame
    tf_sp.word_wrap = True
    for idx, (bld, nrm) in enumerate(spawner_points):
        p = tf_sp.paragraphs[0] if idx == 0 else tf_sp.add_paragraph()
        rb = p.add_run()
        rb.text = f"✔ {bld} "
        rb.font.bold = True
        rb.font.size = Pt(11.5)
        rb.font.color.rgb = TEXT_CREAM
        rn = p.add_run()
        rn.text = nrm
        rn.font.bold = False
        rn.font.size = Pt(11)
        rn.font.color.rgb = TEXT_MUTED
        p.space_after = Pt(8)

    # Right Card: Object Pooling
    create_card(slide7, Inches(6.78), Inches(1.8), Inches(5.75), Inches(5.1), ACCENT_ORANGE)
    tb_pool_title = slide7.shapes.add_textbox(Inches(7.08), Inches(2.0), Inches(5.15), Inches(0.5))
    p_pt = tb_pool_title.text_frame.paragraphs[0]
    p_pt.text = "⚡ KỸ THUẬT TỐI ƯU OBJECT POOLING"
    p_pt.font.bold = True
    p_pt.font.size = Pt(14)
    p_pt.font.color.rgb = ACCENT_ORANGE

    pool_points = [
        ("Vấn đề của cách làm truyền thống:", "Liên tục gọi `Instantiate()` và `Destroy()` khiến bộ nhớ RAM bị phân mảnh, kích hoạt Garbage Collector (GC) gây giật lag hình ảnh."),
        ("Nguyên lý Object Pooling:", "Khởi tạo sẵn 1 danh sách (Pool) đối tượng quả và bom ngay khi nạp Scene. Giữ chúng ở trạng thái ẩn (`SetActive: false`)."),
        ("Tái chế tức thời (Recycle):", "Khi cần quả mới, Spawner lấy từ Pool và kích hoạt lại. Khi quả rơi khỏi màn hình hoặc bị chém xong, tự động thu hồi về Pool."),
        ("Hiệu quả đo lường thực tế:", "Mức tiêu thụ CPU giảm 40%, thời gian dừng do GC đạt xấp xỉ 0ms, tốc độ khung hình duy trì mượt mà 60 FPS.")
    ]
    tb_pl_body = slide7.shapes.add_textbox(Inches(7.08), Inches(2.6), Inches(5.15), Inches(4.1))
    tf_pl = tb_pl_body.text_frame
    tf_pl.word_wrap = True
    for idx, (bld, nrm) in enumerate(pool_points):
        p = tf_pl.paragraphs[0] if idx == 0 else tf_pl.add_paragraph()
        rb = p.add_run()
        rb.text = f"★ {bld} "
        rb.font.bold = True
        rb.font.size = Pt(11.5)
        rb.font.color.rgb = TEXT_CREAM
        rn = p.add_run()
        rn.text = nrm
        rn.font.bold = False
        rn.font.size = Pt(11)
        rn.font.color.rgb = TEXT_MUTED
        p.space_after = Pt(12)

    # ==========================================
    # SLIDE 8: ĐIỂM SỐ, COMBO & MẠNG SỐNG
    # ==========================================
    slide8 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide8)
    add_header(slide8, "Hệ thống tính điểm & sinh tồn", "HỆ THỐNG ĐIỂM SỐ, COMBO STREAK & MẠNG SỐNG", "Tạo động lực cạnh tranh, tăng tính lôi cuốn và kịch tính trong từng ván chơi")

    score_data = [
        ("01", "CƠ CHẾ TÍNH ĐIỂM & COMBO", TEXT_GOLD, [
            ("Điểm chém cơ bản:", "Mỗi quả hoa quả chém trúng cộng trực tiếp +1 Điểm vào quỹ điểm."),
            ("Hệ thống Combo Streak:", "Nếu người chơi chém trúng từ 3 quả trở lên trong cùng 1 nhát quẹt chuột duy nhất:"),
            ("Thưởng điểm cấp số nhân:", "Điểm cộng thêm = Số quả x Hệ số Combo. Kèm theo hiệu ứng chữ COMBO rực rỡ."),
            ("Tăng cảm giác thỏa mãn:", "Khuyến khích người chơi rình rập cơ hội để gom quả chém một lúc thay vì chém lẻ tẻ.")
        ]),
        ("02", "HỆ THỐNG 3 MẠNG SỐNG", ACCENT_RED, [
            ("Mạng sống trực quan:", "Hiển thị bằng 3 biểu tượng Trái Tim (Hearts) màu đỏ ở góc trên màn hình."),
            ("Cơ chế tha bổng (Forgiving):", "Trái cây rơi xuống đáy màn hình không bị trừ mạng — người chơi thoải mái tận hưởng không sợ áp lực."),
            ("Trừng phạt nghiêm khắc:", "Chém trúng bom phát nổ trừ ngay 1 Tim và rung màn hình cảnh báo."),
            ("Điều kiện Thua Cuộc:", "Khi số Tim về 0, trò chơi lập tức dừng lại và kích hoạt màn hình Game Over.")
        ]),
        ("03", "LƯU TRỮ KỶ LỤC (HIGH SCORE)", ACCENT_GREEN, [
            ("Lưu trữ bền vững:", "Sử dụng công nghệ `PlayerPrefs.SetInt(\"HighScore\", score)` của Unity."),
            ("Tự động so sánh kỷ lục:", "Mỗi khi kết thúc ván đấu, hệ thống tự động kiểm tra điểm số vừa đạt được với kỷ lục cũ."),
            ("Vinh danh thành tích:", "Nếu phá vỡ kỷ lục, màn hình Game Over sẽ hiển thị nhãn \"NEW BEST SCORE!\" nổi bật."),
            ("Lưu giữ qua các phiên chơi:", "Điểm cao nhất vẫn được bảo toàn trọn vẹn ngay cả khi tắt game và mở lại.")
        ])
    ]

    for idx, (num, title, col, pts) in enumerate(score_data):
        l = Inches(0.8 + idx * 4.04)
        create_card(slide8, l, Inches(1.8), Inches(3.64), Inches(5.1))

        hb = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l + Inches(0.15), Inches(1.95), Inches(3.34), Inches(0.8))
        hb.fill.solid()
        hb.fill.fore_color.rgb = CARD_HEADER_BG
        hb.line.color.rgb = col
        hb.line.width = Pt(1)

        tb_h = slide8.shapes.add_textbox(l + Inches(0.2), Inches(2.0), Inches(3.24), Inches(0.7))
        p_h = tb_h.text_frame.paragraphs[0]
        p_h.text = f"{num}. {title}"
        p_h.font.bold = True
        p_h.font.size = Pt(12)
        p_h.font.color.rgb = col

        tb_b = slide8.shapes.add_textbox(l + Inches(0.18), Inches(2.9), Inches(3.28), Inches(3.8))
        tf_b = tb_b.text_frame
        tf_b.word_wrap = True
        for p_i, (bld, nrm) in enumerate(pts):
            p = tf_b.paragraphs[0] if p_i == 0 else tf_b.add_paragraph()
            rb = p.add_run()
            rb.text = f"• {bld} "
            rb.font.bold = True
            rb.font.size = Pt(11)
            rb.font.color.rgb = TEXT_CREAM

            rn = p.add_run()
            rn.text = nrm
            rn.font.bold = False
            rn.font.size = Pt(10.5)
            rn.font.color.rgb = TEXT_MUTED
            p.space_after = Pt(8)

    # ==========================================
    # SLIDE 9: THIẾT KẾ GIAO DIỆN & TRẢI NGHIỆM (UI/UX)
    # ==========================================
    slide9 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide9)
    add_header(slide9, "Mỹ thuật & Giao diện", "THIẾT KẾ GIAO DIỆN & TRẢI NGHIỆM NGƯỜI DÙNG (UI/UX)", "Chủ đề võ quán kiếm đạo Nhật Bản mộc mạc, kết hợp hoạt họa trái cây tươi sáng")

    ui_screens = [
        ("MÀN HÌNH CHÍNH (MAIN MENU)", TEXT_GOLD, [
            ("Không gian Võ Quán:", "Hình nền thớt gỗ cổ điển, ánh sáng tự nhiên ấm áp, tạo cảm giác thân thuộc."),
            ("Tiêu đề nổi bật:", "Logo Fruit Ninja 2D phong cách thư pháp sắc nét, ấn tượng ngay khi mở game."),
            ("Nút Play trực quan:", "Nút bấm phong cách bảng gỗ to rõ, kích thích người chơi bấm bắt đầu ngay."),
            ("Hiển thị Kỷ lục:", "Điểm High Score hiển thị ngay ngoài sảnh để tạo động lực vượt qua thử thách.")
        ]),
        ("GIAO DIỆN IN-GAME (GAMEPLAY HUD)", ACCENT_ORANGE, [
            ("Thiết kế tối giản (Minimal):", "Chỉ hiển thị các thông số sống còn: Điểm số hiện tại và 3 Tim sinh mệnh."),
            ("Góc nhìn thông thoáng:", "Đặt UI ở góc trên, dành trọn vẹn 90% diện tích màn hình cho không gian chém quả."),
            ("Phản hồi xúc giác & Âm thanh:", "Âm thanh chém kiếm ngọt lịm khi cắt quả, tiếng xì xì ngòi nổ tạo cảm giác hồi hộp."),
            ("Hiệu ứng té nước ép:", "Vệt nước ép trái cây bám nhẹ tạo cảm giác đã tay đã mắt.")
        ]),
        ("MÀN HÌNH GAME OVER & THỐNG KÊ", ACCENT_GREEN, [
            ("Bảng thông báo chiến tích:", "Hộp thoại tổng kết rõ ràng Điểm số vừa đạt được và Điểm kỷ lục."),
            ("Hiệu ứng New Best Score:", "Huy hiệu vinh danh rực rỡ khi người chơi thiết lập cột mốc kỷ lục mới."),
            ("Tái đấu tức thì (Quick Restart):", "Chỉ cần 1 nút bấm Restart để reset bàn chơi trong 0.1 giây, không để người chơi phải chờ đợi.")
        ])
    ]

    for idx, (title, col, pts) in enumerate(ui_screens):
        l = Inches(0.8 + idx * 4.04)
        create_card(slide9, l, Inches(1.8), Inches(3.64), Inches(5.1))

        hb = slide9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l + Inches(0.15), Inches(1.95), Inches(3.34), Inches(0.8))
        hb.fill.solid()
        hb.fill.fore_color.rgb = CARD_HEADER_BG
        hb.line.color.rgb = col
        hb.line.width = Pt(1)

        tb_h = slide9.shapes.add_textbox(l + Inches(0.2), Inches(2.0), Inches(3.24), Inches(0.7))
        p_h = tb_h.text_frame.paragraphs[0]
        p_h.text = title
        p_h.font.bold = True
        p_h.font.size = Pt(11.5)
        p_h.font.color.rgb = col

        tb_b = slide9.shapes.add_textbox(l + Inches(0.18), Inches(2.9), Inches(3.28), Inches(3.8))
        tf_b = tb_b.text_frame
        tf_b.word_wrap = True
        for p_i, (bld, nrm) in enumerate(pts):
            p = tf_b.paragraphs[0] if p_i == 0 else tf_b.add_paragraph()
            rb = p.add_run()
            rb.text = f"★ {bld} "
            rb.font.bold = True
            rb.font.size = Pt(11)
            rb.font.color.rgb = TEXT_CREAM

            rn = p.add_run()
            rn.text = nrm
            rn.font.bold = False
            rn.font.size = Pt(10.5)
            rn.font.color.rgb = TEXT_MUTED
            p.space_after = Pt(8)

    # ==========================================
    # SLIDE 10: KẾT LUẬN & HƯỚNG PHÁT TRIỂN
    # ==========================================
    slide10 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide10)
    add_header(slide10, "Tổng kết đồ án", "KẾT LUẬN & HƯỚNG PHÁT TRIỂN DỰ ÁN", "Đánh giá mức độ hoàn thiện sản phẩm và định hướng nâng cấp trong tương lai")

    # Left: Results achieved
    create_card(slide10, Inches(0.8), Inches(1.8), Inches(5.75), Inches(3.5), TEXT_GOLD)
    tb_res_t = slide10.shapes.add_textbox(Inches(1.1), Inches(1.95), Inches(5.15), Inches(0.45))
    p_rt = tb_res_t.text_frame.paragraphs[0]
    p_rt.text = "🏆 KẾT QUẢ ĐẠT ĐƯỢC CỦA ĐỒ ÁN"
    p_rt.font.bold = True
    p_rt.font.size = Pt(14)
    p_rt.font.color.rgb = TEXT_GOLD

    res_points = [
        ("Hoàn thiện 100% Core Gameplay:", "Cơ chế chém quả, nảy vật lý, bẫy bom, tính điểm và mạng sống vận hành hoàn hảo."),
        ("Hiệu năng tối ưu tuyệt đối:", "Áp dụng thành công Object Pooling giúp game đạt tốc độ 60+ FPS ổn định, mượt mà."),
        ("Kiến trúc mã nguồn chuẩn mực:", "Tổ chức code sạch sẽ, tuân thủ nguyên lý OOP, Singleton và tách biệt các Manager rõ ràng."),
        ("Đồ họa & Âm thanh sống động:", "Giao diện phong cách võ quán ấm cúng, hiệu ứng vệt kiếm và rung camera cuốn hút.")
    ]
    tb_res_b = slide10.shapes.add_textbox(Inches(1.1), Inches(2.45), Inches(5.15), Inches(2.7))
    tf_rb = tb_res_b.text_frame
    tf_rb.word_wrap = True
    for idx, (bld, nrm) in enumerate(res_points):
        p = tf_rb.paragraphs[0] if idx == 0 else tf_rb.add_paragraph()
        rb = p.add_run()
        rb.text = f"✔ {bld} "
        rb.font.bold = True
        rb.font.size = Pt(11)
        rb.font.color.rgb = TEXT_CREAM
        rn = p.add_run()
        rn.text = nrm
        rn.font.bold = False
        rn.font.size = Pt(10.5)
        rn.font.color.rgb = TEXT_MUTED
        p.space_after = Pt(4)

    # Right: Future roadmap
    create_card(slide10, Inches(6.78), Inches(1.8), Inches(5.75), Inches(3.5), ACCENT_ORANGE)
    tb_fut_t = slide10.shapes.add_textbox(Inches(7.08), Inches(1.95), Inches(5.15), Inches(0.45))
    p_ft = tb_fut_t.text_frame.paragraphs[0]
    p_ft.text = "🚀 HƯỚNG PHÁT TRIỂN TRONG TƯƠNG LAI"
    p_ft.font.bold = True
    p_ft.font.size = Pt(14)
    p_ft.font.color.rgb = ACCENT_ORANGE

    fut_points = [
        ("Thêm Chế Độ Chơi Mới:", "Bổ sung chế độ tính giờ Zen Mode (90s không bom) và Arcade Mode với nhiều loại trái cây ma thuật."),
        ("Trái Cây Kỹ Năng Đặc Biệt:", "Chuối Freeze (đóng băng làm chậm thời gian), Chuối Frenzy (bắn bão hoa quả liên tục)."),
        ("Tùy Biến Lưỡi Kiếm & Dojo:", "Cửa hàng Skin kiếm (Blade Skins) với hiệu ứng vệt kiếm lửa, băng, sấm sét và đổi nền võ quán."),
        ("Bảng Xếp Hạng Trực Tuyến:", "Tích hợp Leaderboard Online (Firebase / PlayFab) để người chơi so tài toàn cầu.")
    ]
    tb_fut_b = slide10.shapes.add_textbox(Inches(7.08), Inches(2.45), Inches(5.15), Inches(2.7))
    tf_fb = tb_fut_b.text_frame
    tf_fb.word_wrap = True
    for idx, (bld, nrm) in enumerate(fut_points):
        p = tf_fb.paragraphs[0] if idx == 0 else tf_fb.add_paragraph()
        rb = p.add_run()
        rb.text = f"✦ {bld} "
        rb.font.bold = True
        rb.font.size = Pt(11)
        rb.font.color.rgb = TEXT_CREAM
        rn = p.add_run()
        rn.text = nrm
        rn.font.bold = False
        rn.font.size = Pt(10.5)
        rn.font.color.rgb = TEXT_MUTED
        p.space_after = Pt(4)

    # Bottom: Thank You & Q&A Banner
    thank_card = create_card(slide10, Inches(0.8), Inches(5.5), Inches(11.733), Inches(1.4), TEXT_GOLD, CARD_HEADER_BG)
    tb_thx = slide10.shapes.add_textbox(Inches(1.0), Inches(5.6), Inches(11.333), Inches(1.2))
    tf_thx = tb_thx.text_frame
    p_t1 = tf_thx.paragraphs[0]
    p_t1.alignment = PP_ALIGN.CENTER
    p_t1.text = "CHÂN THÀNH CẢM ƠN THẦY CÔ VÀ HỘI ĐỒNG ĐÃ LẮNG NGHE!"
    p_t1.font.bold = True
    p_t1.font.size = Pt(18)
    p_t1.font.color.rgb = TEXT_GOLD

    p_t2 = tf_thx.add_paragraph()
    p_t2.alignment = PP_ALIGN.CENTER
    p_t2.text = "Kính mời Quý Thầy Cô và các bạn đặt câu hỏi đóng góp cho đồ án (Q&A)"
    p_t2.font.size = Pt(13)
    p_t2.font.color.rgb = TEXT_CREAM

    output_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "FRUIT_NINJA_2D_Thuyet_Trinh.pptx")
    prs.save(output_path)
    print(f"Presentation saved successfully at: {output_path}")

if __name__ == "__main__":
    create_presentation()
